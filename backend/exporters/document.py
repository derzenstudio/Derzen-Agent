"""
DERZEN - Document writer.

Takes the text a step produced, reads its shape (headings, lists, tables, code)
and writes it out properly for the format you asked for. A PDF gets real
typography, a spreadsheet gets sized columns and wrapped text instead of
crushed cells, a web page gets a stylesheet you can actually reuse.

Formats: md, txt, html, css, js, pdf, docx, xlsx, csv, json.
Heavy libraries are imported only when that format is asked for, so a missing
one never blocks the rest.
"""
from __future__ import annotations

import csv as csvlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import config
import file_manager

FENCE = chr(96) * 3

DARK = "#17160f"
LIGHT = "#f5f2e9"
ACCENT = "#ab3a20"

EXTENSIONS = {
    "md": ".md",
    "txt": ".txt",
    "html": ".html",
    "css": ".css",
    "js": ".js",
    "pdf": ".pdf",
    "docx": ".docx",
    "xlsx": ".xlsx",
    "csv": ".csv",
    "json": ".json",
}


class ExportError(RuntimeError):
    """Raised when a format cannot be written and the reason is worth showing."""


# ── Reading the shape of the text ────────────────────────────────────────────
def parse(text: str) -> List[Dict[str, Any]]:
    """Turn markdown-ish text into a list of blocks."""
    blocks: List[Dict[str, Any]] = []
    lines = (text or "").replace("\r\n", "\n").split("\n")
    i = 0
    para: List[str] = []

    def flush() -> None:
        if para:
            blocks.append({"kind": "para", "text": " ".join(para).strip()})
            para.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith(FENCE):
            flush()
            language = stripped[3:].strip()
            i += 1
            body: List[str] = []
            while i < len(lines) and not lines[i].strip().startswith(FENCE):
                body.append(lines[i])
                i += 1
            i += 1
            blocks.append({"kind": "code", "language": language, "text": "\n".join(body)})
            continue

        if not stripped:
            flush()
            i += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush()
            blocks.append({"kind": "heading", "level": len(heading.group(1)), "text": heading.group(2).strip()})
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].strip()) <= set("|-: "):
            flush()
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2
            rows: List[List[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append({"kind": "table", "header": header, "rows": rows})
            continue

        if re.match(r"^[-*+]\s+", stripped):
            flush()
            items: List[str] = []
            while i < len(lines) and re.match(r"^[-*+]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*+]\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"kind": "bullets", "items": items})
            continue

        if re.match(r"^\d+[.)]\s+", stripped):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\d+[.)]\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+[.)]\s+", "", lines[i].strip()))
                i += 1
            blocks.append({"kind": "numbers", "items": items})
            continue

        if stripped.startswith(">"):
            flush()
            blocks.append({"kind": "quote", "text": stripped.lstrip("> ").strip()})
            i += 1
            continue

        para.append(stripped)
        i += 1

    flush()
    return blocks


def _plain(text: str) -> str:
    """Strip inline markdown so plain formats do not show stray symbols."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", text)
    return text


def _inline_html(text: str) -> str:
    safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    safe = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", safe)
    safe = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<em>\1</em>", safe)
    safe = re.sub(r"`(.+?)`", r"<code>\1</code>", safe)
    safe = re.sub(r"\[(.+?)\]\((https?://.+?)\)", r'<a href="\2">\1</a>', safe)
    return safe


def _slug(value: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9]+", "-", value or "").strip("-").lower()
    return cleaned or "derzen-output"


def _target(folder: str, name: str, extension: str) -> Path:
    stem = _slug(name)
    if not stem.endswith(extension):
        stem = stem + extension
    relative = f"{folder.strip('/') or 'Reports'}/{stem}"
    path = file_manager.validate_path(relative)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _strip_fence(body: str) -> str:
    """Code answers usually arrive wrapped in a fence. Take the inside."""
    match = re.search(FENCE + r"[a-zA-Z]*\n(.*?)" + FENCE, body, re.S)
    return match.group(1).strip() if match else body.strip()


# ── Text based writers ───────────────────────────────────────────────────────
def to_markdown(title: str, blocks: List[Dict[str, Any]]) -> str:
    out = [f"# {title}", "", f"Written {datetime.now().strftime('%d %B %Y, %H:%M')}", ""]
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            out.append("#" * min(6, block["level"] + 1) + " " + block["text"])
        elif kind == "para":
            out.append(block["text"])
        elif kind == "bullets":
            out.extend(f"- {item}" for item in block["items"])
        elif kind == "numbers":
            out.extend(f"{i}. {item}" for i, item in enumerate(block["items"], start=1))
        elif kind == "quote":
            out.append("> " + block["text"])
        elif kind == "code":
            out.append(FENCE + block.get("language", ""))
            out.append(block["text"])
            out.append(FENCE)
        elif kind == "table":
            out.append("| " + " | ".join(block["header"]) + " |")
            out.append("| " + " | ".join("---" for _ in block["header"]) + " |")
            for row in block["rows"]:
                out.append("| " + " | ".join(row) + " |")
        out.append("")
    return "\n".join(out).strip() + "\n"


def to_text(title: str, blocks: List[Dict[str, Any]]) -> str:
    out = [title.upper(), "=" * len(title), ""]
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            out.extend(["", _plain(block["text"]).upper(), "-" * len(block["text"])])
        elif kind == "para":
            out.extend([_plain(block["text"]), ""])
        elif kind == "bullets":
            out.extend(f"  * {_plain(i)}" for i in block["items"])
            out.append("")
        elif kind == "numbers":
            out.extend(f"  {n}. {_plain(i)}" for n, i in enumerate(block["items"], start=1))
            out.append("")
        elif kind == "quote":
            out.extend([f"  \"{_plain(block['text'])}\"", ""])
        elif kind == "code":
            out.extend(["  " + line for line in block["text"].split("\n")] + [""])
        elif kind == "table":
            widths = [
                max(len(str(block["header"][c])), *(len(str(r[c])) for r in block["rows"] if c < len(r)))
                if block["rows"]
                else len(str(block["header"][c]))
                for c in range(len(block["header"]))
            ]
            def line(cells):
                return "  " + "  ".join(str(cells[c]).ljust(widths[c]) for c in range(len(widths)) if c < len(cells))
            out.append(line(block["header"]))
            out.append("  " + "  ".join("-" * w for w in widths))
            out.extend(line(row) for row in block["rows"])
            out.append("")
    return "\n".join(out).strip() + "\n"


PAGE_CSS = """
:root {
  --gelap: %(dark)s;
  --terang: %(light)s;
  --accent: %(accent)s;
}
* { box-sizing: border-box; margin: 0; padding: 0; border-radius: 0; }
body {
  background: var(--terang);
  color: var(--gelap);
  font-family: 'Poppins', system-ui, sans-serif;
  font-weight: 600;
  line-height: 1.7;
  padding: 7rem 2rem 9rem;
}
main { max-width: 860px; margin: 0 auto; }
h1, h2, h3, h4 { font-weight: 800; line-height: 1.1; letter-spacing: -0.02em; }
h1 { font-size: clamp(2.6rem, 7vw, 4.6rem); margin-bottom: 2.5rem; }
h2 { font-size: 1.9rem; margin: 4.5rem 0 1.2rem; border-left: 3px solid var(--accent); padding-left: 1rem; }
h3 { font-size: 1.25rem; margin: 3rem 0 0.8rem; color: var(--accent); }
p { margin-bottom: 1.4rem; max-width: 68ch; }
.stamp { font-size: 0.8rem; letter-spacing: 0.18em; text-transform: uppercase; border: 1px solid var(--gelap); display: inline-block; padding: 0.45rem 0.9rem; margin-bottom: 4rem; }
ul, ol { margin: 0 0 1.8rem 1.4rem; }
li { margin-bottom: 0.7rem; max-width: 66ch; }
blockquote { border-left: 3px solid var(--accent); padding: 0.4rem 0 0.4rem 1.4rem; margin: 2rem 0; font-style: normal; }
code { background: var(--gelap); color: var(--terang); padding: 0.1rem 0.35rem; font-size: 0.9em; }
pre { background: var(--gelap); color: var(--terang); padding: 1.6rem; overflow-x: auto; margin: 2rem 0; border: 3px solid var(--gelap); }
pre code { background: none; padding: 0; }
table { border-collapse: collapse; width: 100%%; margin: 2.5rem 0; border: 3px solid var(--gelap); }
th { background: var(--gelap); color: var(--terang); text-align: left; font-weight: 800; }
th, td { padding: 0.85rem 1rem; border: 1px solid var(--gelap); vertical-align: top; }
tbody tr:nth-child(even) { background: rgba(23, 22, 15, 0.06); }
a { color: var(--accent); text-underline-offset: 3px; }
""" % {"dark": DARK, "light": LIGHT, "accent": ACCENT}


def to_html(title: str, blocks: List[Dict[str, Any]], standalone: bool = True) -> str:
    body: List[str] = [f"<h1>{_inline_html(title)}</h1>"]
    body.append(f'<p class="stamp">{datetime.now().strftime("%d %B %Y")}</p>')
    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            level = min(4, block["level"] + 1)
            body.append(f"<h{level}>{_inline_html(block['text'])}</h{level}>")
        elif kind == "para":
            body.append(f"<p>{_inline_html(block['text'])}</p>")
        elif kind == "bullets":
            items = "".join(f"<li>{_inline_html(i)}</li>" for i in block["items"])
            body.append(f"<ul>{items}</ul>")
        elif kind == "numbers":
            items = "".join(f"<li>{_inline_html(i)}</li>" for i in block["items"])
            body.append(f"<ol>{items}</ol>")
        elif kind == "quote":
            body.append(f"<blockquote>{_inline_html(block['text'])}</blockquote>")
        elif kind == "code":
            escaped = block["text"].replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            body.append(f"<pre><code>{escaped}</code></pre>")
        elif kind == "table":
            head = "".join(f"<th>{_inline_html(c)}</th>" for c in block["header"])
            rows = "".join(
                "<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in row) + "</tr>"
                for row in block["rows"]
            )
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")
    inner = "\n".join(body)
    if not standalone:
        return inner
    return (
        "<!doctype html>\n<html lang=\"en\">\n<head>\n<meta charset=\"utf-8\">\n"
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{_inline_html(title)}</title>\n"
        '<link rel="preconnect" href="https://fonts.googleapis.com">\n'
        '<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>\n'
        '<link href="https://fonts.googleapis.com/css2?family=Poppins:wght@600;800&display=swap" rel="stylesheet">\n'
        f"<style>{PAGE_CSS}</style>\n</head>\n<body>\n<main>\n{inner}\n</main>\n</body>\n</html>\n"
    )


# ── Binary writers ───────────────────────────────────────────────────────────
def _register_poppins():
    """Use Poppins in PDFs when the files are present, otherwise fall back."""
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError:
        return None, None
    folder = config.ALLOWED_BASE / config.SUBDIRS.get("assets", "Assets") / "fonts"
    bold = folder / "Poppins-ExtraBold.ttf"
    body = folder / "Poppins-SemiBold.ttf"
    if bold.exists() and body.exists():
        try:
            pdfmetrics.registerFont(TTFont("PoppinsXB", str(bold)))
            pdfmetrics.registerFont(TTFont("PoppinsSB", str(body)))
            return "PoppinsXB", "PoppinsSB"
        except Exception:  # noqa: BLE001
            pass
    return None, None


def write_pdf(path: Path, title: str, blocks: List[Dict[str, Any]]) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            ListFlowable,
            ListItem,
            PageBreak,
            Paragraph,
            Preformatted,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise ExportError(
            "PDF export needs reportlab. Install it with: pip install reportlab"
        ) from exc

    heading_font, body_font = _register_poppins()
    heading_font = heading_font or "Helvetica-Bold"
    body_font = body_font or "Helvetica"

    dark = colors.HexColor(DARK)
    light = colors.HexColor(LIGHT)
    accent = colors.HexColor(ACCENT)

    styles = {
        "title": ParagraphStyle("title", fontName=heading_font, fontSize=30, leading=34,
                                textColor=dark, spaceAfter=26, alignment=TA_LEFT),
        "stamp": ParagraphStyle("stamp", fontName=body_font, fontSize=8, leading=12,
                                textColor=dark, spaceAfter=34),
        "h2": ParagraphStyle("h2", fontName=heading_font, fontSize=16, leading=20,
                             textColor=dark, spaceBefore=26, spaceAfter=10),
        "h3": ParagraphStyle("h3", fontName=heading_font, fontSize=12, leading=16,
                             textColor=accent, spaceBefore=18, spaceAfter=7),
        "body": ParagraphStyle("body", fontName=body_font, fontSize=10.5, leading=17,
                               textColor=dark, spaceAfter=11),
        "quote": ParagraphStyle("quote", fontName=body_font, fontSize=10.5, leading=17,
                                textColor=dark, leftIndent=14, borderPadding=0, spaceAfter=13),
        "cell": ParagraphStyle("cell", fontName=body_font, fontSize=9, leading=13, textColor=dark),
        "cellhead": ParagraphStyle("cellhead", fontName=heading_font, fontSize=9, leading=13, textColor=light),
    }

    story: List[Any] = [Paragraph(_inline_html(title), styles["title"])]
    story.append(Paragraph(datetime.now().strftime("%d %B %Y, %H:%M").upper(), styles["stamp"]))

    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            story.append(Paragraph(_inline_html(block["text"]), styles["h2" if block["level"] <= 2 else "h3"]))
        elif kind == "para":
            story.append(Paragraph(_inline_html(block["text"]), styles["body"]))
        elif kind in ("bullets", "numbers"):
            items = [ListItem(Paragraph(_inline_html(i), styles["body"]), leftIndent=16)
                     for i in block["items"]]
            story.append(ListFlowable(items, bulletType="1" if kind == "numbers" else "bullet",
                                      bulletColor=accent, start="1" if kind == "numbers" else None))
            story.append(Spacer(1, 8))
        elif kind == "quote":
            story.append(Paragraph(_inline_html(block["text"]), styles["quote"]))
        elif kind == "code":
            story.append(Preformatted(block["text"], ParagraphStyle(
                "code", fontName="Courier", fontSize=8.5, leading=12, textColor=dark)))
            story.append(Spacer(1, 10))
        elif kind == "table":
            header = [Paragraph(_inline_html(c), styles["cellhead"]) for c in block["header"]]
            body_rows = [[Paragraph(_inline_html(c), styles["cell"]) for c in row] for row in block["rows"]]
            table = Table([header] + body_rows, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), dark),
                ("BOX", (0, 0), (-1, -1), 1.6, dark),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, dark),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ]))
            story.append(Spacer(1, 8))
            story.append(table)
            story.append(Spacer(1, 16))

    document = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=24 * mm, rightMargin=24 * mm, topMargin=26 * mm, bottomMargin=26 * mm,
        title=title, author="DERZEN",
    )
    document.build(story)


def write_docx(path: Path, title: str, blocks: List[Dict[str, Any]]) -> None:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError as exc:
        raise ExportError(
            "Word export needs python-docx. Install it with: pip install python-docx"
        ) from exc

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Poppins"
    normal.font.size = Pt(11)

    heading = document.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x17, 0x16, 0x0F)
    document.add_paragraph(datetime.now().strftime("%d %B %Y, %H:%M"))

    for block in blocks:
        kind = block["kind"]
        if kind == "heading":
            document.add_heading(_plain(block["text"]), level=min(4, block["level"] + 1))
        elif kind == "para":
            document.add_paragraph(_plain(block["text"]))
        elif kind == "bullets":
            for item in block["items"]:
                document.add_paragraph(_plain(item), style="List Bullet")
        elif kind == "numbers":
            for item in block["items"]:
                document.add_paragraph(_plain(item), style="List Number")
        elif kind == "quote":
            document.add_paragraph(_plain(block["text"]), style="Intense Quote")
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block["header"]))
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, block["header"]):
                cell.text = _plain(text)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            for row in block["rows"]:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = _plain(text)
    document.save(str(path))


def write_xlsx(path: Path, title: str, rows: List[List[str]]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ExportError("Spreadsheet export needs openpyxl.") from exc

    if not rows:
        rows = [["No data"], [""]]

    book = Workbook()
    sheet = book.active
    sheet.title = (title or "DERZEN")[:28] or "DERZEN"

    edge = Side(style="thin", color="17160F")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    header_fill = PatternFill("solid", fgColor="17160F")
    band_fill = PatternFill("solid", fgColor="EDE8DA")

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(row=r, column=c, value=value)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if r == 1:
                cell.fill = header_fill
                cell.font = Font(name="Poppins", bold=True, color="F5F2E9", size=11)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                cell.font = Font(name="Poppins", size=10)
                if r % 2 == 1:
                    cell.fill = band_fill

    widths: Dict[int, int] = {}
    for row in rows:
        for c, value in enumerate(row, start=1):
            longest = max((len(part) for part in str(value).split("\n")), default=0)
            widths[c] = min(58, max(widths.get(c, 12), longest + 4))
    for c, width in widths.items():
        sheet.column_dimensions[get_column_letter(c)].width = width

    sheet.row_dimensions[1].height = 26
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    book.save(str(path))


# ── Tables from loose text ───────────────────────────────────────────────────
def _table_from_markdown(text: str) -> List[List[str]]:
    for block in parse(text):
        if block["kind"] == "table":
            return [block["header"]] + block["rows"]
    return []


def _table_from_lines(text: str) -> List[List[str]]:
    rows: List[List[str]] = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if "\t" in line:
            rows.append([c.strip() for c in line.split("\t")])
        elif line.count("|") >= 2:
            rows.append([c.strip() for c in line.strip("|").split("|")])
        elif line.count(",") >= 1 and len(line) < 400:
            rows.append([c.strip() for c in line.split(",")])
        else:
            rows.append([line])
    width = max((len(r) for r in rows), default=1)
    return [r + [""] * (width - len(r)) for r in rows]


async def rows_from_text(text: str) -> List[List[str]]:
    """Ask the local AI for a clean table, fall back to parsing what is there."""
    direct = _table_from_markdown(text)
    if direct:
        return direct

    import ai_manager

    try:
        raw = await ai_manager.query(
            "Turn the material below into one table. Reply with JSON only: an array "
            "of arrays, the first array being the column headings. Six columns at "
            "most. No commentary.\n\n" + text[:9000]
        )
        start, end = raw.find("["), raw.rfind("]")
        parsed = json.loads(raw[start : end + 1])
        rows = [[str(cell) for cell in row] for row in parsed if isinstance(row, list)]
        if rows:
            width = max(len(r) for r in rows)
            return [r + [""] * (width - len(r)) for r in rows]
    except Exception:  # noqa: BLE001 - fall through to the plain parser
        pass
    return _table_from_lines(text)


def save_sheet(filename: str, rows: List[List[str]], folder: str = "Reports", title: str = "") -> dict:
    path = _target(folder, filename, ".xlsx")
    write_xlsx(path, title or filename, rows)
    return {
        "path": str(path.relative_to(config.ALLOWED_BASE)),
        "absolute": str(path),
        "name": path.name,
        "format": "xlsx",
        "rows": len(rows),
    }


# ── The one entry point ──────────────────────────────────────────────────────
def export(
    body: str,
    filename: str,
    fmt: str = "md",
    folder: str = "Reports",
    title: str = "",
    structured: Optional[Any] = None,
) -> dict:
    """Write 'body' to disk in 'fmt' and return where it landed."""
    fmt = (fmt or "md").strip().lower().lstrip(".")
    if fmt not in EXTENSIONS:
        raise ExportError(f"Unknown format '{fmt}'. Pick one of: {', '.join(sorted(EXTENSIONS))}.")

    title = title or filename or "DERZEN output"
    if isinstance(structured, list) and structured:
        blocks = structured
    else:
        source = body if body.strip() else (structured if isinstance(structured, str) else "")
        blocks = parse(source or "")

    path = _target(folder, filename, EXTENSIONS[fmt])

    if fmt == "md":
        path.write_text(to_markdown(title, blocks), encoding="utf-8")
    elif fmt == "txt":
        path.write_text(to_text(title, blocks), encoding="utf-8")
    elif fmt == "html":
        path.write_text(to_html(title, blocks), encoding="utf-8")
    elif fmt == "css":
        content = _strip_fence(body)
        path.write_text(content if content.strip() else PAGE_CSS, encoding="utf-8")
    elif fmt == "js":
        path.write_text(_strip_fence(body) + "\n", encoding="utf-8")
    elif fmt == "json":
        try:
            parsed = json.loads(_strip_fence(body))
        except (json.JSONDecodeError, ValueError):
            parsed = {"title": title, "written": datetime.now().isoformat(timespec="seconds"), "blocks": blocks}
        path.write_text(json.dumps(parsed, indent=2, ensure_ascii=False), encoding="utf-8")
    elif fmt == "csv":
        rows = _table_from_markdown(body) or _table_from_lines(body)
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            csvlib.writer(handle).writerows(rows)
    elif fmt == "xlsx":
        rows = _table_from_markdown(body) or _table_from_lines(body)
        write_xlsx(path, title, rows)
    elif fmt == "pdf":
        write_pdf(path, title, blocks)
    elif fmt == "docx":
        write_docx(path, title, blocks)

    return {
        "path": str(path.relative_to(config.ALLOWED_BASE)),
        "absolute": str(path),
        "name": path.name,
        "format": fmt,
        "bytes": path.stat().st_size if path.exists() else 0,
    }
